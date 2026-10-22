"""Strict tabular import, explicit cleaning and spreadsheet-safe export."""
from io import BytesIO, StringIO
from itertools import chain
import csv
import json
import re
import pandas as pd
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from pathlib import Path
from openpyxl.utils import get_column_letter
from .safety import ValidationError, validate_upload, check_workbook

MAX_ROWS = 100_000
MAX_COLUMNS = 200

def columns(names):
    out = [re.sub(r'[^\w]+', '_', str(n).strip().lower()).strip('_') or f'column_{i+1}' for i,n in enumerate(names)]
    if len(set(out)) != len(out):
        raise ValidationError('Column names collide after standardization. Rename them first.')
    return out

def _frame(headers, rows):
    if not headers or len(headers) > MAX_COLUMNS:
        raise ValidationError('Expected 1–200 columns.')
    headers = [str(x) if x is not None else '' for x in headers]
    if any(not x.strip() for x in headers) or len(set(headers)) != len(headers):
        raise ValidationError('Column headers must be nonempty and unique.')
    if len(rows) > MAX_ROWS:
        raise ValidationError('Maximum 100,000 rows per dataset.')
    return pd.DataFrame(rows, columns=headers)

def read_table(name: str, data: bytes, sheet: str | None = None) -> pd.DataFrame:
    suffix = validate_upload(name, data, {'.csv','.xlsx','.json'})
    try:
        if suffix == '.csv':
            text = data.decode('utf-8-sig')
            reader = csv.reader(StringIO(text), strict=True)
            header = next(reader)
            rows = []
            for row in reader:
                if not row: row = [''] * len(header)
                if len(row) != len(header):
                    raise ValidationError('CSV rows have different column counts. Use comma-separated UTF-8 CSV.')
                rows.append(row)
                if len(rows) > MAX_ROWS: raise ValidationError('Maximum 100,000 rows.')
            # Preserve identifiers/leading zeros; blank strings become missing values.
            return _frame(header, rows).replace('', pd.NA)
        if suffix == '.json':
            def unique(pairs):
                d = {}
                for key,value in pairs:
                    if key in d: raise ValidationError('Duplicate JSON object key: ' + key)
                    d[key] = value
                return d
            def bad_constant(_): raise ValidationError('JSON must not contain NaN or Infinity.')
            records = json.loads(data.decode('utf-8-sig'), object_pairs_hook=unique, parse_constant=bad_constant)
            if not isinstance(records,list) or not records or not all(isinstance(r,dict) and r for r in records):
                raise ValidationError('JSON must be a nonempty array of flat objects.')
            keys = list(records[0])
            if any(set(r) != set(keys) for r in records):
                raise ValidationError('JSON records must have identical keys; normalize missing keys first.')
            if any(isinstance(v,(dict,list)) for r in records for v in r.values()):
                raise ValidationError('Nested JSON is unsupported. Flatten it explicitly first.')
            return _frame(keys, [[r[k] for k in keys] for r in records])
        check_workbook(data)
        wb = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
        try:
            if sheet is not None and sheet not in wb.sheetnames:
                raise ValidationError('Worksheet not found.')
            ws = wb[sheet] if sheet else wb.worksheets[0]
            if (ws.max_row or 0)>MAX_ROWS+1 or (ws.max_column or 0)>MAX_COLUMNS:
                raise ValidationError('Workbook exceeds 100,000 rows or 200 columns.')
            # Read actual cells, including workbooks without dimensions or with false metadata.
            ws.reset_dimensions()
            iterator=ws.iter_rows(values_only=True)
            header=next(iterator,None)
            if not header:raise ValidationError('Worksheet is empty.')
            if len(header)>MAX_COLUMNS:raise ValidationError('Maximum 200 columns.')
            rows=[]
            for row in iterator:
                if len(row)>len(header):raise ValidationError('Worksheet has data beyond its header columns.')
                rows.append(list(row)+[None]*(len(header)-len(row)))
                if len(rows)>MAX_ROWS:raise ValidationError('Maximum 100,000 rows.')
            return _frame(header,rows)
        finally: wb.close()
    except ValidationError: raise
    except Exception as e:
        raise ValidationError('Could not read file. Check its format, encoding and structure.') from e

def sheet_names(name, data):
    validate_upload(name, data, {'.xlsx'})
    check_workbook(data)
    try:
        wb=load_workbook(BytesIO(data),read_only=True,keep_links=False)
        try:return wb.sheetnames
        finally:wb.close()
    except Exception as e:raise ValidationError('Cannot read workbook worksheets.') from e

def merge_tables(frames):
    if not frames: raise ValidationError('No tables selected.')
    expected=list(frames[0].columns)
    if any(set(f.columns)!=set(expected) for f in frames):
        raise ValidationError('Merge requires identical column names. Standardize headers or select compatible files.')
    if sum(len(f) for f in frames)>MAX_ROWS:raise ValidationError('Merged result exceeds 100,000 rows.')
    return pd.concat([f.reindex(columns=expected) for f in frames],ignore_index=True)

def profile(df):
    return {'rows':len(df),'columns':len(df.columns),'duplicate_rows':int(df.duplicated().sum()),
            'missing_cells':int(df.isna().sum().sum()),'missing_by_column':df.isna().sum().astype(int).to_dict()}

def clean_table(df, *, trim=True, text_case='unchanged', normalize_headers=True, remove_empty=True, deduplicate=True):
    out=df.copy(deep=True)
    changed=0
    for c in out.columns:
        def transform(v):
            nonlocal changed
            if not isinstance(v,str):return v
            new=v.strip() if trim else v
            if text_case=='lower':new=new.lower()
            elif text_case=='upper':new=new.upper()
            elif text_case=='title':new=new.title()
            elif text_case!='unchanged':raise ValidationError('Unknown text formatting option.')
            if new!=v:changed+=1
            return pd.NA if new=='' else new
        out[c]=out[c].map(transform)
    if normalize_headers:out.columns=columns(out.columns)
    empty=int(out.isna().all(axis=1).sum()) if remove_empty else 0
    if remove_empty:out=out.dropna(how='all')
    duplicates=int(out.duplicated().sum()) if deduplicate else 0
    if deduplicate:out=out.drop_duplicates()
    out=out.reset_index(drop=True)
    return out,{'rows_before':len(df),'rows_after':len(out),'empty_rows_removed':empty,
                'duplicate_rows_removed':duplicates,'text_cells_modified':changed,
                'headers_changed':sum(a!=b for a,b in zip(df.columns,out.columns)),
                'missing_values_after':int(out.isna().sum().sum())}

def filter_sort(df, column=None, operator='contains', value='', sort_by=None, ascending=True):
    out=df.copy()
    if column:
        if column not in out:raise ValidationError('Filter column not found.')
        s=out[column]
        if operator=='contains':mask=s.astype('string').str.contains(value,regex=False,case=False,na=False)
        elif operator=='equals':mask=s.astype('string').eq(value).fillna(False)
        elif operator=='missing':mask=s.isna()
        elif operator in {'greater than','less than'}:
            try:threshold=float(value)
            except ValueError as e:raise ValidationError('Enter a numeric filter value.') from e
            nums=pd.to_numeric(s,errors='coerce')
            mask=nums>threshold if operator=='greater than' else nums<threshold
        else:raise ValidationError('Unknown filter operation.')
        out=out.loc[mask]
    if sort_by:
        if sort_by not in out:raise ValidationError('Sort column not found.')
        try:out=out.sort_values(sort_by,ascending=ascending,kind='stable',na_position='last')
        except TypeError as e:raise ValidationError('Sort column mixes incompatible data types.') from e
    return out.reset_index(drop=True)

def _safe_cell(v):
    if isinstance(v,str) and v.lstrip().startswith(('=','+','-','@','\t','\r')):return "'"+v
    return v

def export_csv(df):
    out=df.map(_safe_cell)
    out.columns=[_safe_cell(c) for c in out.columns]
    return out.to_csv(index=False,lineterminator='\n').encode('utf-8-sig')

def export_excel(df):
    if len(df)>MAX_ROWS or len(df.columns)>MAX_COLUMNS:
        raise ValidationError('Export exceeds dataset limit.')
    # Validate before OpenPyXL starts its temporary worksheet stream.
    for row in chain([list(df.columns)],df.itertuples(index=False,name=None)):
        for value in row:
            if isinstance(value,str):
                if len(value)>32767:
                    raise ValidationError('Excel cells cannot exceed 32,767 characters; export CSV instead.')
                if ILLEGAL_CHARACTERS_RE.search(value):
                    raise ValidationError('Excel cannot store these control characters; export CSV instead.')
    wb=Workbook(write_only=True)
    ws=wb.create_sheet('Results')
    ws.freeze_panes='A2'
    ws.auto_filter.ref=f'A1:{get_column_letter(max(1,len(df.columns)))}{len(df)+1}'
    try:
        for index,row in enumerate(chain([list(df.columns)],df.itertuples(index=False,name=None))):
            cooked=[]
            for value in row:
                if pd.isna(value):value=None
                cell=WriteOnlyCell(ws,value=value)
                if isinstance(value,str):cell.data_type='s'
                if index==0:
                    cell.font=Font(bold=True,color='FFFFFF')
                    cell.fill=PatternFill('solid',fgColor='167D8D')
                cooked.append(cell)
            ws.append(cooked)
        buf=BytesIO()
        wb.save(buf)
        return buf.getvalue()
    except Exception as error:
        raise ValidationError('Excel export failed. Try CSV or a smaller dataset.') from error
    finally:
        # OpenPyXL uses a secure temporary worksheet file internally. Its normal
        # save path deletes it, but failed saves need explicit cleanup as well.
        # _writer is pinned-version integration, covered by cleanup regression tests.
        try:
            if not ws.closed:
                ws.close()
        finally:
            writer=getattr(ws,'_writer',None)
            if writer is not None and Path(writer.out).exists():
                writer.cleanup()
            wb.close()
