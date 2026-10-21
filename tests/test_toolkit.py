from io import BytesIO
from pathlib import Path
import json
import zipfile
import pandas as pd
import pytest
from openpyxl import load_workbook
from toolkit.safety import ValidationError, safe_name, validate_upload, validate_batch, check_workbook, MAX_FILE
from toolkit.tabular import read_table, merge_tables, clean_table, columns, filter_sort, export_csv, export_excel, sheet_names
from toolkit.files import plan_files, build_zip, FileOperation
from toolkit.pdf import extract_pdf
from generate_samples import generate

ROOT=Path(__file__).resolve().parents[1]

@pytest.fixture(scope='session',autouse=True)
def samples():generate()

@pytest.mark.parametrize('name',['../x','/tmp/x','a\\b','a/b','..','bad\x00name','CON.txt','x.','x:evil',''])
def test_unsafe_names(name):
    with pytest.raises(ValidationError):safe_name(name)

def test_upload_limits():
    with pytest.raises(ValidationError):validate_upload('x.csv',b'')
    with pytest.raises(ValidationError):validate_upload('x.csv',b'x'*(MAX_FILE+1))
    with pytest.raises(ValidationError):validate_batch([('x.txt',b'x')]*31)
    with pytest.raises(ValidationError):validate_upload('x.exe',b'x',{'.csv'})

@pytest.mark.parametrize('raw',[b'id,name\n1,A,B',b'foo,foo\n1,2',b',name\n1,A',b'\xff\xfe',b'"broken'])
def test_bad_csv(raw):
    with pytest.raises(ValidationError):read_table('x.csv',raw)

@pytest.mark.parametrize('raw',[b'{}',b'[]',b'[1]',b'[{"a":{}}]',b'[{"a":1},{"b":2}]',b'[{"a":1,"a":2}]',b'[{"a":NaN}]',b'not json'])
def test_bad_json(raw):
    with pytest.raises(ValidationError):read_table('x.json',raw)

def test_clean_merge_and_preserve_original():
    a=read_table('a.csv',(ROOT/'samples/input/sales_january.csv').read_bytes())
    b=read_table('b.xlsx',(ROOT/'samples/input/sales_february.xlsx').read_bytes())
    a.columns=columns(a.columns);b.columns=columns(b.columns)
    original=merge_tables([a,b]);copy=original.copy(deep=True)
    result,summary=clean_table(original)
    pd.testing.assert_frame_equal(original,copy)
    assert len(result)==4 and summary['duplicate_rows_removed']==1 and summary['empty_rows_removed']==1
    assert result.iloc[0]['order_id']=='001'
    assert result.iloc[0]['customer_name']=='Alice Chen'
    assert result['region'].isna().sum()==1
    assert len(filter_sort(result,'amount','greater than','100'))==2
    assert len(filter_sort(result,'region','equals','North'))==2
    assert len(filter_sort(result,'region','missing'))==1
    assert filter_sort(result,sort_by='order_id',ascending=False).iloc[0]['order_id']=='004'

def test_incompatible_and_collision():
    with pytest.raises(ValidationError):merge_tables([pd.DataFrame(columns=['a']),pd.DataFrame(columns=['b'])])
    with pytest.raises(ValidationError):columns(['Customer Name','customer-name'])
    with pytest.raises(ValidationError):filter_sort(pd.DataFrame({'n':[1]}),'n','greater than','abc')

def test_reordered_merge():
    result=merge_tables([pd.DataFrame({'a':[1],'b':[2]}),pd.DataFrame({'b':[4],'a':[3]})])
    assert result.to_dict('records')==[{'a':1,'b':2},{'a':3,'b':4}]

def test_text_options():
    result,s=clean_table(pd.DataFrame({' Name ':[' alice ','ALICE',' ']}),text_case='upper')
    assert result.to_dict('list')=={'name':['ALICE']}
    assert s['text_cells_modified']==2

def test_export_formulas_and_roundtrip():
    frame=pd.DataFrame({'id':['001','=1+1',' +SUM(A1)','@danger']})
    data=export_excel(frame)
    wb=load_workbook(BytesIO(data));assert wb.active['A3'].data_type=='s';wb.close()
    assert read_table('x.xlsx',data).iloc[1,0]=='=1+1'
    assert "'=1+1" in export_csv(frame).decode('utf-8-sig')
    assert sheet_names('x.xlsx',data)==['Results']
    assert read_table('x.xlsx',data).iloc[0,0]=='001'

def test_bad_excel_and_zip_traversal():
    with pytest.raises(ValidationError):read_table('x.xlsx',b'not excel')
    buf=BytesIO()
    with zipfile.ZipFile(buf,'w') as z:
        z.writestr('[Content_Types].xml','x');z.writestr('xl/workbook.xml','x');z.writestr('../escape','x')
    with pytest.raises(ValidationError):check_workbook(buf.getvalue())

def test_archive_preview_exact_bytes():
    files=[('sales.csv',b'a,b\n1,2'),('invoice.pdf',b'original bytes')]
    plan=plan_files(files,prefix='client_',numbered=True,organize='extension')
    with zipfile.ZipFile(BytesIO(build_zip(files,plan))) as z:
        assert z.namelist()==['csv/001_client_sales.csv','pdf/002_client_invoice.pdf']
        assert z.read(z.namelist()[0])==files[0][1]
    assert files[0][0]=='sales.csv'

def test_archive_collision_paths_and_rules():
    files=[('a.txt',b'a'),('A.txt',b'b')]
    with pytest.raises(ValidationError):plan_files(files)
    with pytest.raises(ValidationError):plan_files([('x.txt',b'a')],prefix='../')
    good=[('sales.csv',b'a')]
    plan=plan_files(good,organize='prefix rules',rules_text='sales=reports')
    assert plan[0].destination=='reports/sales.csv'
    with pytest.raises(ValidationError):build_zip(good,[FileOperation(0,'sales.csv','../escape',1)])
    with pytest.raises(ValidationError):build_zip(good,[])

def test_pdf_actual_table_and_scanned():
    result=extract_pdf('invoice.pdf',(ROOT/'samples/input/fictional_invoice.pdf').read_bytes())
    assert 'INV-001' in result.pages.iloc[0]['text']
    assert len(result.tables)==1
    assert result.tables[0][1].iloc[1,0]=='Data cleanup'
    scan=extract_pdf('scan.pdf',(ROOT/'samples/input/scanned_invoice.pdf').read_bytes())
    assert scan.pages.iloc[0]['text']=='' and not scan.tables
    assert any('OCR' in w for w in scan.warnings)
    with pytest.raises(ValidationError):extract_pdf('bad.pdf',b'%PDF-broken')

def test_sample_outputs():
    result=read_table('cleaned.csv',(ROOT/'samples/output/cleaned_sales.csv').read_bytes())
    assert len(result)==4
    summary=json.loads((ROOT/'samples/output/pdf_extraction_summary.json').read_text())
    assert summary['tables']==1 and summary['pages']==1

def test_header_formula_protection():
    frame=pd.DataFrame({'=HYPERLINK("bad")':['001']})
    assert "'=HYPERLINK" in export_csv(frame).decode('utf-8-sig')
    wb=load_workbook(BytesIO(export_excel(frame)))
    assert wb.active['A1'].data_type=='s'
    wb.close()

def test_long_excel_text_has_explicit_error():
    frame=pd.DataFrame({'text':['x'*32768]})
    with pytest.raises(ValidationError):export_excel(frame)
    assert len(export_csv(frame))>32768

def test_pdf_encryption_and_page_limit():
    import pymupdf
    doc=pymupdf.open();doc.new_page()
    raw=doc.tobytes(encryption=pymupdf.PDF_ENCRYPT_AES_256,user_pw='secret',owner_pw='owner')
    with pytest.raises(ValidationError,match='Password'):extract_pdf('protected.pdf',raw)
    for _ in range(100):doc.new_page()
    with pytest.raises(ValidationError,match='100'):extract_pdf('long.pdf',doc.tobytes())
    doc.close()

def test_workbook_expansion_limit():
    buf=BytesIO()
    with zipfile.ZipFile(buf,'w',compression=zipfile.ZIP_DEFLATED) as z:
        z.writestr('[Content_Types].xml','x');z.writestr('xl/workbook.xml','x')
        z.writestr('xl/oversized.xml','x'*2_000_000)
    with pytest.raises(ValidationError,match='compression ratio'):check_workbook(buf.getvalue())

def test_duplicate_workbook_members_rejected():
    import warnings
    buf=BytesIO()
    with warnings.catch_warnings():
        warnings.simplefilter('ignore',UserWarning)
        with zipfile.ZipFile(buf,'w') as z:
            z.writestr('[Content_Types].xml','x')
            z.writestr('xl/workbook.xml','x')
            z.writestr('xl/workbook.xml','replacement')
    with pytest.raises(ValidationError,match='duplicate'):check_workbook(buf.getvalue())

def test_excel_xml_entity_is_rejected():
    from openpyxl.xml import DEFUSEDXML
    assert DEFUSEDXML, 'DefusedXML must be active for Excel parsing.'
    source=export_excel(pd.DataFrame({'id':['001']}))
    xml=b'<?xml version="1.0"?><!DOCTYPE x [<!ENTITY a "blocked">]><worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData><row r="1"><c r="A1" t="inlineStr"><is><t>&a;</t></is></c></row></sheetData></worksheet>'
    buf=BytesIO()
    with zipfile.ZipFile(BytesIO(source)) as original, zipfile.ZipFile(buf,'w') as target:
        for member in original.infolist():
            target.writestr(member.filename,xml if member.filename=='xl/worksheets/sheet1.xml' else original.read(member))
    with pytest.raises(ValidationError):read_table('entity.xlsx',buf.getvalue())
